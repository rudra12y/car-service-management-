import os
import openpyxl
from flask import Flask, request, jsonify
from flask_cors import CORS

app = Flask(__name__)

# Allow all origins for CORS to fix frontend & browser connection issues
CORS(app, resources={r"/*": {"origins": "*"}})

EXCEL_FILE = "car_service_orders.xlsx"

def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Bookings"
        ws.append(["Name", "Phone", "Car Model", "Service Date", "Service Type", "Status"])
        wb.save(EXCEL_FILE)

init_excel()

@app.route("/", methods=["GET"])
def home():
    return "Server is Active", 200

@app.route("/api/bookings", methods=["GET"])
def get_bookings():
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
        bookings = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(row):
                bookings.append({
                    "name": row[0],
                    "phone": row[1],
                    "car_model": row[2],
                    "service_date": str(row[3]),
                    "service_type": row[4],
                    "status": row[5] if len(row) > 5 else "Pending"
                })
        return jsonify(bookings), 200
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/bookings", methods=["POST"])
def add_booking():
    try:
        data = request.json
        if not data:
            return jsonify({"error": "No data received"}), 400

        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
        ws.append([
            data.get("name"),
            data.get("phone"),
            data.get("car_model"),
            data.get("service_date"),
            data.get("service_type"),
            "Pending"
        ])
        wb.save(EXCEL_FILE)
        return jsonify({"message": "Booking successful"}), 201
    except Exception as e:
        return jsonify({"error": str(e)}), 500

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000)